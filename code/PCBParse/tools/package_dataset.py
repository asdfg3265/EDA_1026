#!/usr/bin/env python3
"""Package a validated dataset for the data-contribution bonus.

Usage:
    python tools/package_dataset.py dataset/ --team ABC123 [--desc my_description.docx]

Produces team_<team>_result_dataset.zip containing
    dataset/<case_id>.png, dataset/<case_id>.json
    dataset_description.docx  (generated if --desc is not given)
"""
import os
import sys
import zipfile
import json
import argparse

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from validate_dataset import check_case, CASE_RE  # noqa: E402


# The 34 component types of the official type_list (10GT-0930/type_list.md; the 10-22
# evaluation note's appendix): crystals and op-amps have no pin-count variants there.
TYPE_LIST = {'gnd', 'c', 'r', 'v', 'box', 'net_input', 'net_short', 'net_output', 'net_bidirection', 'd', 'led',
             'pin', 'switch', 'l', 'mosfet', 'other', 'm3螺丝', 'bjt', 'testpoint', 'crystal', 'buzzer', 'battary',
             'fuse', 'block', 'amp', 'jumper', 'motor', 'esd', 'dc', 'transfomer', 'ant', 'circle_header', 'opt', 'seg'}
TYPE_MAP = {'crystal_2pin': 'crystal', 'crystal_3pin': 'crystal', 'crystal_4pin': 'crystal',
            'amp_3pin': 'amp', 'amp_5pin': 'amp', 'amp_7pin': 'amp',
            'bjt_npn': 'bjt', 'bjt_pnp': 'bjt', 'mosfet_pnp': 'mosfet', 'mosfet_npn': 'mosfet'}


def to_type_list(J):
    """Map pin-count / polarity sub-types onto the official type list; returns the
    number of changed components and the set of types still outside the list."""
    n = 0
    for v in J.get('components', {}).values():
        t = v.get('type')
        if t in TYPE_MAP:
            v['type'] = TYPE_MAP[t]
            n += 1
    bad = {v.get('type') for v in J.get('components', {}).values()} - TYPE_LIST
    return n, bad


def _png_size(path):
    import struct
    with open(path, 'rb') as f:
        h = f.read(24)
    return struct.unpack('>II', h[16:24])


def drop_offimage(J, W, H):
    """Drop parts drawn outside the rendered page (mounting holes placed off the sheet):
    nothing of them is visible, so they cannot be labelled.  Only parts with no net
    connections are dropped; returns their keys."""
    used = set()
    for n in (J.get('nets') or {}).values():
        hg = n.get('hyperGraph', '') if isinstance(n, dict) else ''
        used.update(tok.split('.')[0] for tok in str(hg).replace('(', ' ').replace(')', ' ').replace(',', ' ').split())
    out = []
    for k, v in list((J.get('components') or {}).items()):
        b = v.get('bbox') or []
        if len(b) != 4:
            continue
        inside_w = max(0, min(b[2], W) - max(b[0], 0))
        inside_h = max(0, min(b[3], H) - max(b[1], 0))
        area = max(1, (b[2] - b[0]) * (b[3] - b[1]))
        if inside_w * inside_h < 0.5 * area and k not in used:
            del J['components'][k]
            (J.get('pins') or {}).pop(k, None)
            out.append(k)
    return out


def make_description(path, team, n_cases, stats):
    try:
        import docx
    except ImportError:
        raise SystemExit('python-docx is required to generate the description (pip install python-docx), '
                         'or pass --desc with your own document')
    d = docx.Document()
    d.add_heading('数据说明文档', 0)
    d.add_paragraph(f'队伍标识：{team}')
    d.add_heading('1. 数据来源', 1)
    d.add_paragraph('原理图来自嘉立创开源硬件平台（oshwhub.com）上以开源许可证公开的嘉立创 EDA 专业版工程。'
                    '只使用作者公开了源文件的工程；作者未公开源文件的工程一律跳过。'
                    '许可证禁止演绎的工程（CC BY-ND、CC BY-NC-ND 等）已排除。'
                    '每个案例对应的工程网址、作者与许可证列在压缩包内的 sources.csv 中；'
                    '使用 ShareAlike / NonCommercial 许可的案例请按原许可证使用（仅用于学术研究与开源数据集）。'
                    '已与赛题提供的 200 个训练案例按元件签名去重；同一工程最多取 8 张图并去除近似重复的版本。')
    d.add_heading('2. 数量', 1)
    d.add_paragraph(f'共 {n_cases} 个案例，每个案例包含同名 PNG 原图与 JSON 标注（与赛题输出格式一致）。')
    d.add_heading('3. 图片生成方式', 1)
    d.add_paragraph('图片由本队编写的渲染程序从工程源文件（原理图页 + 符号库）直接绘制：按嘉立创 EDA 专业版编辑器的默认主题'
                    '（背景 #FFFCF8、导线 #008800、符号 #A00000/#880000、文字颜色与字号按源文件设置），'
                    '用 Chromium 栅格化，比例约 1.9–2.05 像素/单位，去除图框与标题栏，四周留白。'
                    '我们将渲染结果与编辑器导出的 SVG 逐元素比对过：引脚位置一致率 99% 以上，连接点 100% 一致。')
    d.add_heading('4. 标注方法', 1)
    d.add_paragraph('标注不是人工绘制的，而是由转换程序从同一份源文件计算得到：元件位号、类型、型号/取值、引脚号与引脚名、'
                    '引脚端点坐标、网络连接（hyperGraph，按导线与引脚的几何连通性计算，不按同名标签合并）和导线线段'
                    '（在连接点处拆分，每条线段恰好两个端点），坐标与图片像素一一对应。'
                    'components 以位号为键；电源/地符号以 ∅+元素编号为键；网络标签以标签文字为键，重名按元素顺序加 _2、_3 后缀；'
                    'pins 键为 pin_<引脚号>。转换规则与赛题公开的 golden case 逐项对照确定。'
                    '全部案例都通过了自动检查（Schema、引用一致性、整数像素坐标、有效性：≥1 个 IC 且有效引脚数达标、'
                    '≥5 个元件、≥5 个有效网络），并抽样生成标注叠加图进行人工目视检查。')
    d.add_heading('5. 坐标系声明', 1)
    d.add_paragraph('所有绝对坐标为图内像素坐标，以图片左下角为原点，x 轴向右增大，y 轴向上增大，'
                    '一律为整数像素（粒度 1 像素）；元件框 bbox 为 [x1, y1, x2, y2]，(x1, y1) 为左下角，(x2, y2) 为右上角。'
                    '不使用相对坐标。')
    d.add_heading('6. 统计', 1)
    for k, v in stats.items():
        d.add_paragraph(f'{k}: {v}')
    d.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('folder')
    ap.add_argument('--team', required=True)
    ap.add_argument('--desc', default=None)
    ap.add_argument('--out', default='.')
    ap.add_argument('--sources', default=None, help='cases.csv from build_dataset (project url, licence, author)')
    ap.add_argument('--exclude-licence', default=r'\bND\b|NoDeriv', help='regex of licences to leave out')
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    if not CASE_RE.match(a.team):
        raise SystemExit('team id must use ASCII letters, digits and "_" only')
    import csv
    import re as _re
    src = {}
    srcp = a.sources or os.path.join(a.folder, 'cases.csv')
    if os.path.exists(srcp):
        for r in csv.DictReader(open(srcp, encoding='utf-8-sig')):
            src[r['case_id']] = r
    else:
        print('WARNING: no cases.csv - sources.csv will be missing')
    exl = _re.compile(a.exclude_licence, _re.I)
    cases = sorted(os.path.splitext(f)[0] for f in os.listdir(a.folder)
                   if f.lower().endswith('.json') and f.startswith('case_'))
    good = []
    n_mapped = 0
    tot_c = tot_n = 0
    for cid in cases:
        png = os.path.join(a.folder, cid + '.png')
        if not os.path.exists(png):
            print('skip (no png):', cid)
            continue
        if cid in src and exl.search(src[cid].get('license') or ''):
            print('skip (licence):', cid, src[cid].get('license'))
            continue
        if src and (src.get(cid, {}).get('license') or '').strip().lower() in ('', 'null', 'none'):
            print('skip (no licence stated):', cid)
            continue
        J = json.load(open(os.path.join(a.folder, cid + '.json'), encoding='utf-8'))
        dropped = drop_offimage(J, *_png_size(png))
        if dropped:
            print('dropped (outside the page):', cid, dropped)
        err, warn, st = check_case(cid, J)
        if err:
            print('skip (invalid):', cid, err[:3])
            continue
        nmap, bad = to_type_list(J)
        if bad:
            print('skip (type outside type_list):', cid, sorted(bad))
            continue
        n_mapped += nmap
        good.append(cid)
        tot_c += st['components']
        tot_n += st['nets']
    if len(good) < 20:
        print(f'WARNING: only {len(good)} valid cases (>= 20 required)')
    desc = a.desc
    if desc is None:
        desc = os.path.join(a.out, 'dataset_description.docx')
        make_description(desc, a.team, len(good), {'有效案例数': len(good), '元件总数': tot_c, '有效网络总数': tot_n})
    zpath = os.path.join(a.out, f'team_{a.team}_result_dataset.zip')
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for cid in good:
            z.write(os.path.join(a.folder, cid + '.png'), f'dataset/{cid}.png')
            J = json.load(open(os.path.join(a.folder, cid + '.json'), encoding='utf-8'))
            drop_offimage(J, *_png_size(os.path.join(a.folder, cid + '.png')))
            to_type_list(J)
            z.writestr(f'dataset/{cid}.json', json.dumps(J, ensure_ascii=False, indent=1))
        ext = os.path.splitext(desc)[1].lower()
        z.write(desc, 'dataset_description' + ext)
        if src:
            import io
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(['case_id', 'project_url', 'author', 'license', 'sheet'])
            for cid in good:
                r = src.get(cid, {})
                w.writerow([cid, r.get('project'), r.get('author'), r.get('license'), r.get('sheet')])
            z.writestr('sources.csv', '\ufeff' + buf.getvalue())
    print('types mapped onto the official type_list:', n_mapped, 'components')
    print('written', zpath, 'with', len(good), 'cases')


if __name__ == '__main__':
    main()
